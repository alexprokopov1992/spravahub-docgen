from pathlib import Path
import re
import sys

from docx import Document
from starlette.datastructures import FormData

# Allows running as: python scripts/validate_templates.py
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.forms import build_values  # noqa: E402

PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def paragraphs_in(container):
    for paragraph in getattr(container, "paragraphs", []):
        yield paragraph
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from paragraphs_in(cell)


def all_text(doc: Document) -> str:
    parts = []
    parts.extend(p.text for p in paragraphs_in(doc))
    for section in doc.sections:
        for container in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            parts.extend(p.text for p in paragraphs_in(container))
    return "\n".join(parts)


def main() -> int:
    project = Path(__file__).resolve().parents[1]
    template_dir = project / "local_templates"
    supported, _ = build_values(FormData({}))
    supported_keys = set(supported)

    files = [
        template_dir / "application_recalc.docx",
        template_dir / "claim_indexation.docx",
        template_dir / "claim_recalc.docx",
    ]

    failed = False
    for path in files:
        if not path.exists():
            print(f"ERROR {path.name}: file not found")
            failed = True
            continue
        doc = Document(str(path))
        markers = {m.group(1).strip() for m in PLACEHOLDER_RE.finditer(all_text(doc))}
        unsupported = sorted(markers - supported_keys)
        if unsupported:
            print(f"ERROR {path.name}: unsupported placeholders: {', '.join(unsupported)}")
            failed = True
        else:
            print(f"OK {path.name}: {len(markers)} placeholders, all supported")

    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
