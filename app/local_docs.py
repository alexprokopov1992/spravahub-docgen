from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Dict
import re
import subprocess

from docx import Document
from docx.oxml.ns import qn

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"
PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


@dataclass
class GeneratedFile:
    data: bytes
    filename: str
    media_type: str


def _paragraphs_in_container(container):
    for p in getattr(container, "paragraphs", []):
        yield p
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from _paragraphs_in_container(cell)


def _all_paragraphs(doc: Document):
    yield from _paragraphs_in_container(doc)
    for section in doc.sections:
        for container in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            yield from _paragraphs_in_container(container)


def _document_text(doc: Document) -> str:
    return "\n".join(p.text for p in _all_paragraphs(doc))


def _markers(doc: Document) -> set[str]:
    return {m.group(1).strip() for m in PLACEHOLDER_RE.finditer(_document_text(doc))}




def _clear_red_highlight(run) -> None:
    """Remove only the red authoring highlight used to mark template variables."""
    rpr = run._r.rPr
    if rpr is None:
        return

    # Google Docs red background is exported to DOCX as w:highlight w:val="red".
    for highlight in list(rpr.findall(qn("w:highlight"))):
        if (highlight.get(qn("w:val")) or "").lower() == "red":
            rpr.remove(highlight)

    # Be defensive for DOCX masters edited outside Google Docs.
    for shading in list(rpr.findall(qn("w:shd"))):
        fill = (shading.get(qn("w:fill")) or "").replace("#", "").upper()
        color = (shading.get(qn("w:color")) or "").replace("#", "").upper()
        if fill in {"FF0000", "F00"} or color in {"FF0000", "F00"}:
            rpr.remove(shading)


def _clear_template_red_highlights(doc: Document) -> None:
    """Strip red variable-marker highlighting, preserving all other formatting."""
    for paragraph in _all_paragraphs(doc):
        for run in paragraph.runs:
            _clear_red_highlight(run)

        # Google Docs can also export highlight on the paragraph mark itself.
        ppr = paragraph._p.pPr
        if ppr is not None:
            paragraph_rpr = ppr.find(qn("w:rPr"))
            if paragraph_rpr is not None:
                for highlight in list(paragraph_rpr.findall(qn("w:highlight"))):
                    if (highlight.get(qn("w:val")) or "").lower() == "red":
                        paragraph_rpr.remove(highlight)
                for shading in list(paragraph_rpr.findall(qn("w:shd"))):
                    fill = (shading.get(qn("w:fill")) or "").replace("#", "").upper()
                    color = (shading.get(qn("w:color")) or "").replace("#", "").upper()
                    if fill in {"FF0000", "F00"} or color in {"FF0000", "F00"}:
                        paragraph_rpr.remove(shading)


def _replace_across_runs(paragraph, marker: str, replacement: str) -> int:
    runs = list(paragraph.runs)
    if not runs:
        return 0

    count = 0
    while True:
        texts = [r.text or "" for r in runs]
        full = "".join(texts)
        start = full.find(marker)
        if start < 0:
            break
        end = start + len(marker)

        bounds = []
        pos = 0
        for idx, text in enumerate(texts):
            bounds.append((idx, pos, pos + len(text)))
            pos += len(text)

        start_info = next((b for b in bounds if b[1] <= start < b[2]), None)
        end_pos = max(start, end - 1)
        end_info = next((b for b in bounds if b[1] <= end_pos < b[2]), None)
        if start_info is None or end_info is None:
            break

        si, ss, _ = start_info
        ei, es, _ = end_info
        prefix = runs[si].text[: start - ss]
        suffix = runs[ei].text[end - es :]

        if si == ei:
            runs[si].text = prefix + replacement + suffix
        else:
            # Keep the formatting of the run where the placeholder starts.
            runs[si].text = prefix + replacement
            for i in range(si + 1, ei):
                runs[i].text = ""
            runs[ei].text = suffix
        count += 1
    return count


def render_docx(template_path: Path, values: Dict[str, str], output_path: Path):
    doc = Document(str(template_path))
    template_markers = _markers(doc)
    unsupported = sorted(template_markers.difference(values.keys()))
    if unsupported:
        raise RuntimeError(
            "Template contains unsupported placeholders: " + ", ".join(unsupported)
        )

    # Replace exactly the placeholders that are actually present in this template.
    paragraphs = list(_all_paragraphs(doc))
    for key in sorted(template_markers, key=len, reverse=True):
        marker = "{{" + key + "}}"
        value = values.get(key, "") or ""
        for paragraph in paragraphs:
            _replace_across_runs(paragraph, marker, value)

    unresolved = sorted(_markers(doc))
    if unresolved:
        raise RuntimeError(
            "Generated document still contains placeholders: " + ", ".join(unresolved)
        )

    # Red highlight is an authoring aid in the master template, not output content.
    _clear_template_red_highlights(doc)

    doc.save(str(output_path))


def convert_to_pdf(docx_path: Path, pdf_dir: Path) -> Path:
    pdf_dir.mkdir(parents=True, exist_ok=True)
    profile = pdf_dir / "lo-profile"
    profile.mkdir(exist_ok=True)
    cmd = [
        "libreoffice",
        "-env:UserInstallation=file://" + str(profile.resolve()),
        "--headless",
        "--convert-to", "pdf",
        "--outdir", str(pdf_dir),
        str(docx_path),
    ]
    cp = subprocess.run(cmd, capture_output=True, text=True, timeout=90)
    pdf_path = pdf_dir / (docx_path.stem + ".pdf")
    if cp.returncode != 0 or not pdf_path.exists():
        raise RuntimeError(
            "LibreOffice PDF conversion failed: "
            + (cp.stderr.strip() or cp.stdout.strip() or "unknown error")
        )
    return pdf_path


class LocalDocsRenderer:
    def __init__(self, template_dir: str, pdf_enabled: bool = True):
        self.template_dir = Path(template_dir)
        self.pdf_enabled = pdf_enabled

    @staticmethod
    def _safe_name(name: str) -> str:
        return (
            "".join(ch if ch.isalnum() or ch in " _-." else "_" for ch in name)
            .strip(" ._") or "document"
        )

    def render(
        self,
        template_filename: str,
        values: Dict[str, str],
        output_format: str,
        download_basename: str,
    ) -> GeneratedFile:
        template_path = self.template_dir / template_filename
        if not template_path.exists():
            raise RuntimeError(
                f"Local template not found: {template_path}. "
                "Run scripts/sync_local_templates.py."
            )

        safe = self._safe_name(download_basename)
        with TemporaryDirectory(dir="/srv/app/tmp") as td:
            td = Path(td)
            docx_path = td / f"{safe}.docx"
            render_docx(template_path, values, docx_path)

            if output_format == "docx":
                return GeneratedFile(
                    data=docx_path.read_bytes(),
                    filename=docx_path.name,
                    media_type=DOCX_MIME,
                )

            if output_format == "pdf":
                if not self.pdf_enabled:
                    raise RuntimeError("Local PDF generation is disabled.")
                pdf_path = convert_to_pdf(docx_path, td)
                return GeneratedFile(
                    data=pdf_path.read_bytes(),
                    filename=pdf_path.name,
                    media_type=PDF_MIME,
                )

            raise ValueError("Unsupported output format.")
